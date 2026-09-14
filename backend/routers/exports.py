import io
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from database import supabase
from security import get_current_user

router = APIRouter(prefix="/api/export", tags=["Exports"])

BIMONTHLY_LABELS = {
    1: 'January - February', 2: 'January - February',
    3: 'March - April', 4: 'March - April',
    5: 'May - June', 6: 'May - June',
    7: 'July - August', 8: 'July - August',
    9: 'September - October', 10: 'September - October',
    11: 'November - December', 12: 'November - December'
}

def get_period_name(m: int) -> str:
    return BIMONTHLY_LABELS.get(m, f"Period {m}")

@router.get("/excel/{bill_id}")
def export_excel(bill_id: str, user: dict = Depends(get_current_user)):
    b = supabase.table("monthly_bills").select("*, entries:bill_entries(*, house:houses(*)), society:societies(*)").eq("id", bill_id).execute()
    if not b.data:
        raise HTTPException(status_code=404, detail="Bill not found")

    bill = b.data[0]
    entries = bill.get("entries") or []
    entries.sort(key=lambda x: x.get("house", {}).get("houseNo", ""))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Water Bill"

    # Header
    title = f"{bill.get('society', {}).get('name', 'Society')} - Water Bill ({get_period_name(bill['month'])} {bill['year']})"
    ws.merge_cells("A1:G1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["Home No", "H.V", "A.V", "UNIT", "V", "FALO", "TOTAL"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")

    for col in range(1, 8):
        cell = ws.cell(row=2, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row_idx = 3
    for e in entries:
        h_no = e.get("house", {}).get("houseNo", "")
        ws.append([
            h_no,
            e.get("hv"),
            e.get("av"),
            e.get("unit"),
            e.get("v", 600),
            e.get("falo"),
            e.get("total")
        ])
        row_idx += 1

    # Totals Row
    ws.append(["TOTALS", "", "", f"=SUM(D3:D{row_idx-1})", f"=SUM(E3:E{row_idx-1})", f"=SUM(F3:F{row_idx-1})", f"=SUM(G3:G{row_idx-1})"])
    for col in range(1, 8):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Water_Bill_{get_period_name(bill['month'])}_{bill['year']}.xlsx".replace(" ", "_")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/pdf/{bill_id}")
def export_pdf(bill_id: str, user: dict = Depends(get_current_user)):
    b = supabase.table("monthly_bills").select("*, entries:bill_entries(*, house:houses(*)), society:societies(*)").eq("id", bill_id).execute()
    if not b.data:
        raise HTTPException(status_code=404, detail="Bill not found")

    bill = b.data[0]
    entries = bill.get("entries") or []
    entries.sort(key=lambda x: x.get("house", {}).get("houseNo", ""))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(name='TitleStyle', parent=styles['Heading1'], fontSize=16, leading=20, textColor=colors.HexColor('#1E3A8A'), alignment=1)
    story.append(Paragraph(f"<b>{bill.get('society', {}).get('name', 'Society')}</b>", title_style))
    story.append(Paragraph(f"Water Bill Report — {get_period_name(bill['month'])} {bill['year']}", ParagraphStyle(name='Sub', alignment=1, fontSize=11, textColor=colors.HexColor('#4B5563'))))
    story.append(Spacer(1, 15))

    data = [["Home No", "H.V", "A.V", "UNIT", "V", "FALO", "TOTAL"]]
    for e in entries:
        data.append([
            str(e.get("house", {}).get("houseNo", "")),
            str(e.get("hv") if e.get("hv") is not None else "—"),
            str(e.get("av") if e.get("av") is not None else "—"),
            str(e.get("unit") if e.get("unit") is not None else "—"),
            str(e.get("v", 600)),
            str(e.get("falo") if e.get("falo") is not None else "—"),
            str(e.get("total") if e.get("total") is not None else "—")
        ])

    table = Table(data, colWidths=[60, 75, 75, 75, 60, 75, 80])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')])
    ]))

    story.append(table)
    doc.build(story)
    buffer.seek(0)

    filename = f"Water_Bill_{get_period_name(bill['month'])}_{bill['year']}.pdf".replace(" ", "_")
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
